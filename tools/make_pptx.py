"""Writes an unlabeled .pptx the way PowerPoint lays one out (python-pptx).

Each slide is (title, body lines, table rows); `notes` become the first
slide's speaker notes. Covers the places a deck keeps text: title and body
placeholders, a table, and the notes page.
"""
from pptx import Presentation
from pptx.util import Inches

TITLE_AND_CONTENT = 1
TITLE_ONLY = 5


def write_pptx(path: str, slides: list[tuple[str, list[str], list[list[str]]]], notes: str = ""):
    deck = Presentation()
    for index, (title, body, table_rows) in enumerate(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[TITLE_AND_CONTENT if body else TITLE_ONLY])
        slide.shapes.title.text = title
        if body:
            slide.placeholders[1].text = "\n".join(body)
        if table_rows:
            table = slide.shapes.add_table(len(table_rows), len(table_rows[0]),
                                           Inches(1), Inches(2), Inches(8), Inches(1)).table
            for row, cells in enumerate(table_rows):
                for column, value in enumerate(cells):
                    table.cell(row, column).text = value
        if index == 0 and notes:
            slide.notes_slide.notes_text_frame.text = notes
    deck.save(path)
